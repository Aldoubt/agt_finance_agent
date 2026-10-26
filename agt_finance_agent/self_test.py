"""Offline smoke check used by Windows packaging, never by the normal UI."""

from __future__ import annotations

import json
from pathlib import Path
import sys
import tempfile
import traceback


def run_self_test(workspace_root=None, *, check_ocr: bool = False) -> int:
    root = Path(workspace_root) if workspace_root else Path(tempfile.mkdtemp(prefix="agt-finance-smoke-"))
    root.mkdir(parents=True, exist_ok=True)
    report = {"success": False, "frozen": bool(getattr(sys, "frozen", False)), "ocr_checked": False}
    try:
        from agt_finance_agent import __version__
        from core.workbench.batch import BatchSession
        from core.generators.excel_generator import ExcelGenerator
        from openpyxl import load_workbook

        session = BatchSession.create_demo(root / "batches")
        for invoice in session.invoices:
            session.confirm_invoice(invoice["id"])
        for asset in session.assets:
            if asset["kind"] in {"photo", "purchase_proof", "payment_proof"}:
                selected = asset.get("suggested_item_ids") or asset.get("item_ids") or []
                session.set_links(asset["id"], selected, confirmed=True)
        output = session.export(root / "output")
        assert len(list(output.glob("*.pdf"))) == 2
        assert len(list(output.iterdir())) == 4
        wb = load_workbook(output / "采购统计.xlsx", data_only=True)
        ws = wb["采购统计"]
        assert [cell.value for cell in ws[1]] == ExcelGenerator.HEADERS
        assert ws.max_column == 8
        assert sum(float(ws.cell(row, 3).value) for row in (2, 3, 4)) == 258.0
        assert ws.cell(ws.max_row, 3).value == 258.0
        assert ws.cell(2, 8).value == ws.cell(3, 8).value
        wb.close()

        if check_ocr:
            from core.parser.screenshot_parser import ScreenshotParser
            sample = next(asset for asset in session.assets if asset["kind"] == "purchase_proof")
            parsed = ScreenshotParser().parse(session.asset_path(sample))
            report["ocr_checked"] = True
            report["ocr_text_characters"] = len(parsed.text)

        report.update(success=True, version=__version__, invoice_count=2, item_count=3,
                      total_amount="258.00", excel_columns=8, output=str(output))
        exit_code = 0
    except Exception as exc:
        report.update(error=str(exc), traceback=traceback.format_exc())
        exit_code = 1
    (root / "self_test_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    if sys.stdout is not None:
        print(json.dumps(report, ensure_ascii=False))
    return exit_code
