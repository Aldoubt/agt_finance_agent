"""Validate a confirmed batch once and generate its three coordinated outputs.

The plan contains confirmed values, never a second parsing of the source invoice.
Original PDFs are copied byte for byte.  Draft errors are available to the UI via
ValidationError.preview; only a valid plan may be exported.
"""

from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
import hashlib
import os
from pathlib import Path, PureWindowsPath
import re
import shutil
import tempfile
from typing import Callable
import uuid

import fitz
from PIL import Image
from docx import Document
from openpyxl import load_workbook

from .docx_generator import DocxGenerator
from .excel_generator import ExcelGenerator


CENT = Decimal("0.01")
EVIDENCE_KINDS = ("photo", "purchase_proof", "payment_proof")


class ValidationError(ValueError):
    """Actionable validation errors, with a partial preview for unfinished drafts."""

    def __init__(self, errors, warnings=(), preview=None):
        self.errors = list(errors)
        self.warnings = list(warnings)
        self.preview = preview or {}
        super().__init__("\n".join(str(row.get("message", row)) for row in self.errors))


def _problem(code, message, **context):
    return {"code": code, "message": message, **context}


def _money(value, errors, *, label, **context):
    if value is None or isinstance(value, bool) or str(value).strip() == "":
        errors.append(_problem("missing_amount", f"{label}还没有金额，请核对后填写。", **context))
        return None
    try:
        number = Decimal(str(value).strip())
        if not number.is_finite():
            raise InvalidOperation
        cents = number.quantize(CENT)
        if cents != number:
            errors.append(_problem(
                "amount_precision", f"{label}的金额必须精确到分，不能包含超过两位有效小数。",
                **context,
            ))
            return None
        # Avoid a negative-zero folder name while retaining legitimate zero amounts.
        return abs(cents) if cents == 0 else cents
    except (InvalidOperation, ValueError, TypeError):
        errors.append(_problem("invalid_amount", f"{label}的金额不是有效数字。", **context))
        return None


def _money_text(value):
    return None if value is None else format(value, ".2f")


def _safe_name(value, max_length=54):
    value = re.sub(r'[\x00-\x1f<>:"/\\|?*]+', "_", str(value))
    value = re.sub(r"\s+", " ", value).strip(" .")
    return (value or "商品")[:max_length].rstrip(" .") or "商品"


def _sha256(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _batch_file(batch_path, relative_path):
    if not isinstance(relative_path, str) or not relative_path:
        raise ValueError("材料没有有效的保存路径。")
    raw = Path(relative_path)
    windows = PureWindowsPath(relative_path)
    if raw.is_absolute() or windows.drive or ".." in raw.parts or ".." in windows.parts:
        raise ValueError("材料路径必须位于当前批次内。")
    root = Path(batch_path).resolve()
    resolved = (root / raw).resolve()
    try:
        resolved.relative_to(root)
    except ValueError as exc:
        raise ValueError("材料路径不能指向当前批次之外。") from exc
    if not resolved.is_file():
        raise FileNotFoundError("材料文件不存在，请重新添加。")
    return resolved


@dataclass(frozen=True)
class SourceAsset:
    id: str
    relative_path: str
    source: Path
    sha256: str
    kind: str
    original_name: str


def _check_source(source, batch_path, *, inspect_content=True):
    errors = []
    label = source.original_name or source.id
    try:
        path = _batch_file(batch_path, source.relative_path)
        if path != source.source:
            raise ValueError("材料路径在复核后发生变化，请重新添加。")
        if _sha256(path) != source.sha256:
            errors.append(_problem(
                "source_changed", f"材料“{label}”在导入后发生变化，请重新添加并核对。",
                asset_id=source.id,
            ))
            return errors
        if inspect_content and source.kind == "invoice":
            with fitz.open(path) as pdf:
                if not pdf.is_pdf or pdf.needs_pass or len(pdf) == 0:
                    raise ValueError("发票不是可读取的 PDF，或需要密码。")
        elif inspect_content and source.kind in EVIDENCE_KINDS:
            with Image.open(path) as picture:
                picture.verify()
    except (OSError, ValueError, RuntimeError) as exc:
        errors.append(_problem(
            "invalid_source", f"无法使用材料“{label}”：{exc}", asset_id=source.id,
        ))
    return errors


@dataclass(frozen=True)
class InvoiceExport:
    id: str
    asset_id: str
    sequence: int
    filename: str
    source: Path | None
    total: Decimal | None
    invoice_no: str
    supplier: str
    item_ids: tuple[str, ...]


@dataclass(frozen=True)
class ItemExport:
    id: str
    sequence: int
    invoice_id: str
    invoice_sequence: int
    name: str
    raw_name: str
    amount: Decimal | None
    invoice_filename: str
    invoice_source: Path | None
    photo_ids: tuple[str, ...]
    purchase_proof_ids: tuple[str, ...]
    payment_proof_ids: tuple[str, ...]
    photo_files: tuple[str, ...]
    purchase_proof_files: tuple[str, ...]
    payment_files: tuple[str, ...]

    @property
    def gross_amount(self):
        return self.amount

    @property
    def invoice_files(self):
        return () if self.invoice_source is None else (str(self.invoice_source),)


@dataclass(frozen=True)
class ExportPlan:
    batch_id: str
    batch_path: Path
    invoices: tuple[InvoiceExport, ...]
    items: tuple[ItemExport, ...]
    sources: tuple[SourceAsset, ...]
    total: Decimal | None
    warnings: tuple[dict, ...]

    @property
    def invoice_exports(self):
        return [
            {
                "id": row.id, "asset_id": row.asset_id, "sequence": row.sequence,
                "source": str(row.source) if row.source else None,
                "filename": row.filename, "total": _money_text(row.total),
                "invoice_no": row.invoice_no, "supplier": row.supplier,
                "item_ids": list(row.item_ids),
            }
            for row in self.invoices
        ]

    @property
    def excel_rows(self):
        return [
            [
                item.sequence, item.name,
                float(item.amount) if item.amount is not None else None,
                "是" if item.invoice_source else "否",
                "是" if item.photo_files else "否",
                "是" if item.purchase_proof_files else "否",
                "是" if item.payment_files else "否",
                item.invoice_filename,
            ]
            for item in self.items
        ]

    @property
    def counts(self):
        return {
            "invoices": len(self.invoices),
            "items": len(self.items),
            "photos": len({p for item in self.items for p in item.photo_ids}),
            "purchase_proofs": len({p for item in self.items for p in item.purchase_proof_ids}),
            "payment_proofs": len({p for item in self.items for p in item.payment_proof_ids}),
        }

    def to_dict(self):
        return {
            "batch_id": self.batch_id,
            "invoices": self.invoice_exports,
            "items": [
                {
                    "id": item.id, "sequence": item.sequence,
                    "invoice_id": item.invoice_id, "invoice_sequence": item.invoice_sequence,
                    "name": item.name, "raw_name": item.raw_name,
                    "amount": _money_text(item.amount),
                    "invoice_filename": item.invoice_filename,
                    "photo_ids": list(item.photo_ids),
                    "purchase_proof_ids": list(item.purchase_proof_ids),
                    "payment_proof_ids": list(item.payment_proof_ids),
                    "photo_files": list(item.photo_files),
                    "purchase_proof_files": list(item.purchase_proof_files),
                    "payment_files": list(item.payment_files),
                }
                for item in self.items
            ],
            "excel_headers": list(ExcelGenerator.HEADERS),
            "excel_rows": self.excel_rows,
            "total": _money_text(self.total),
            "warnings": list(self.warnings),
            "counts": self.counts,
        }


def _records(data, key, errors):
    value = data.get(key, [])
    if not isinstance(value, list):
        errors.append(_problem("invalid_batch", f"{key} 数据格式不正确。"))
        return []
    rows = []
    seen = set()
    for row in value:
        if not isinstance(row, dict) or not isinstance(row.get("id"), str) or not row["id"]:
            errors.append(_problem("invalid_batch", f"{key} 中有缺少编号的记录。"))
            continue
        if row["id"] in seen:
            errors.append(_problem("duplicate_id", f"{key} 中存在重复记录编号。", record_id=row["id"]))
            continue
        seen.add(row["id"])
        rows.append(row)
    return rows


def build_export_plan(data: dict, batch_path: Path) -> ExportPlan:
    """Build a stable export snapshot, or raise with errors and a partial preview."""
    errors, warnings = [], []
    root = Path(batch_path).resolve()
    assets_list = _records(data, "assets", errors)
    all_invoices = _records(data, "invoices", errors)
    all_items = _records(data, "items", errors)
    assets = {row["id"]: row for row in assets_list}
    invoices = [
        row for row in all_invoices
        if not assets.get(row.get("asset_id"), {}).get("ignored", False)
    ]
    active_invoice_ids = {row["id"] for row in invoices}
    all_invoice_ids = {row["id"] for row in all_invoices}
    items = [row for row in all_items if row.get("invoice_id") in active_invoice_ids]
    for item in all_items:
        if item.get("invoice_id") not in all_invoice_ids:
            errors.append(_problem(
                "unknown_invoice", f"商品“{item.get('name', '')}”没有对应发票。",
                item_id=item["id"],
            ))
    item_ids = {row["id"] for row in items}
    all_item_ids = {row["id"] for row in all_items}
    if not invoices:
        errors.append(_problem("no_invoice", "请先添加至少一张发票，照片和整理进度可以先保存。"))
    source_map = {}

    def use_source(asset_id, expected_kind):
        asset = assets.get(asset_id)
        if asset is None:
            errors.append(_problem("missing_asset", "关联的材料记录不存在，请重新添加。", asset_id=asset_id))
            return None
        if asset.get("kind") != expected_kind:
            errors.append(_problem(
                "wrong_material_type", f"材料“{asset.get('original_name', '')}”的类型与关联不符。",
                asset_id=asset_id,
            ))
            return None
        if asset_id in source_map:
            return source_map[asset_id]
        expected_hash = str(asset.get("sha256", "")).lower()
        if not re.fullmatch(r"[0-9a-f]{64}", expected_hash):
            errors.append(_problem("missing_source_hash", "材料校验信息不完整，请重新添加。", asset_id=asset_id))
            return None
        try:
            path = _batch_file(root, asset.get("path"))
        except (OSError, ValueError) as exc:
            errors.append(_problem(
                "invalid_source", f"材料“{asset.get('original_name', '')}”不可用：{exc}",
                asset_id=asset_id,
            ))
            return None
        source = SourceAsset(
            id=asset_id, relative_path=asset["path"], source=path,
            sha256=expected_hash, kind=expected_kind,
            original_name=str(asset.get("original_name") or path.name),
        )
        source_map[asset_id] = source
        errors.extend(_check_source(source, root))
        return source

    links = {item["id"]: {kind: [] for kind in EVIDENCE_KINDS} for item in items}
    for asset in assets_list:
        kind = asset.get("kind")
        if asset.get("ignored") or kind not in EVIDENCE_KINDS:
            continue
        linked_ids = asset.get("item_ids", [])
        if not isinstance(linked_ids, list):
            errors.append(_problem("invalid_relation", "材料关联格式不正确。", asset_id=asset["id"]))
            continue
        if asset.get("confirmed") is not True:
            warnings.append(_problem(
                "unconfirmed_material",
                f"材料“{asset.get('original_name', '')}”还没有确认关联，本次不会放入 Word。",
                asset_id=asset["id"],
            ))
            continue
        unknown = [value for value in linked_ids if value not in all_item_ids]
        if unknown:
            errors.append(_problem(
                "unknown_item", "材料关联到了不存在的商品，请重新选择。", asset_id=asset["id"],
            ))
        active_links = [value for value in linked_ids if value in item_ids]
        if not active_links:
            continue
        source = use_source(asset["id"], kind)
        if source:
            for item_id in dict.fromkeys(active_links):
                links[item_id][kind].append(source.id)

    invoice_exports, item_exports = [], []
    invoice_numbers, invoice_hashes = {}, {}
    width = max(2, len(str(len(invoices))))
    for sequence, invoice in enumerate(invoices, 1):
        invoice_id = invoice["id"]
        label = f"第 {sequence} 张发票"
        if invoice.get("confirmed") is not True:
            errors.append(_problem(
                "unconfirmed_invoice", f"{label}还没有核对确认。", invoice_id=invoice_id,
            ))
        invoice_total = _money(
            invoice.get("total"), errors, label=label, invoice_id=invoice_id,
        )
        source = use_source(invoice.get("asset_id"), "invoice")
        invoice_no = re.sub(r"\s+", "", str(invoice.get("invoice_no") or ""))
        if invoice_no:
            if invoice_no in invoice_numbers:
                errors.append(_problem(
                    "duplicate_invoice", f"发票号码 {invoice_no} 重复，请忽略重复导入的那一张。",
                    invoice_id=invoice_id, other_invoice_id=invoice_numbers[invoice_no],
                ))
            invoice_numbers[invoice_no] = invoice_id
        else:
            warnings.append(_problem(
                "missing_invoice_number", f"{label}没有填写发票号码，请确认原件并非重复导入。",
                invoice_id=invoice_id,
            ))
        if source:
            if source.sha256 in invoice_hashes:
                errors.append(_problem(
                    "duplicate_invoice_file", f"{label}与另一张发票内容完全相同，请忽略重复文件。",
                    invoice_id=invoice_id, other_invoice_id=invoice_hashes[source.sha256],
                ))
            invoice_hashes[source.sha256] = invoice_id
        invoice_items = [row for row in items if row.get("invoice_id") == invoice_id]
        if not invoice_items:
            errors.append(_problem(
                "missing_items", f"{label}还没有商品明细，请按发票补充商品和含税金额。",
                invoice_id=invoice_id,
            ))
        names = [str(row.get("name") or row.get("raw_name") or "").strip() for row in invoice_items]
        amount_label = _money_text(invoice_total) or "待核对"
        filename = f"{sequence:0{width}d}_{_safe_name('+'.join(names))}_{amount_label}元.pdf"
        amounts = []
        for item, name in zip(invoice_items, names):
            item_id = item["id"]
            if not name:
                errors.append(_problem("missing_item_name", "请填写商品名称。", item_id=item_id))
            if re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", name):
                errors.append(_problem("invalid_item_name", "商品名称包含不可打印字符，请重新填写。", item_id=item_id))
            amount = _money(item.get("amount"), errors, label=f"商品“{name}”", item_id=item_id)
            amounts.append(amount)
            relation = links[item_id]
            for kind, message in (
                ("photo", "实物图"), ("purchase_proof", "购买凭证"),
            ):
                if not relation[kind]:
                    warnings.append(_problem(
                        f"missing_{kind}", f"商品“{name}”没有关联{message}，可稍后补充。",
                        item_id=item_id,
                    ))
            item_exports.append(ItemExport(
                id=item_id, sequence=len(item_exports) + 1, invoice_id=invoice_id,
                invoice_sequence=sequence, name=name,
                raw_name=str(item.get("raw_name") or name), amount=amount,
                invoice_filename=filename, invoice_source=source.source if source else None,
                photo_ids=tuple(relation["photo"]),
                purchase_proof_ids=tuple(relation["purchase_proof"]),
                payment_proof_ids=tuple(relation["payment_proof"]),
                photo_files=tuple(str(source_map[x].source) for x in relation["photo"]),
                purchase_proof_files=tuple(str(source_map[x].source) for x in relation["purchase_proof"]),
                payment_files=tuple(str(source_map[x].source) for x in relation["payment_proof"]),
            ))
        if invoice_total is not None and amounts and all(value is not None for value in amounts):
            item_total = sum(amounts, Decimal("0.00"))
            if item_total != invoice_total:
                errors.append(_problem(
                    "invoice_item_amount_mismatch",
                    f"{label}的商品含税金额合计 {item_total:.2f} 元，与发票金额 {invoice_total:.2f} 元不一致。",
                    invoice_id=invoice_id,
                ))
        invoice_exports.append(InvoiceExport(
            id=invoice_id, asset_id=str(invoice.get("asset_id") or ""), sequence=sequence,
            filename=filename, source=source.source if source else None,
            total=invoice_total, invoice_no=invoice_no, supplier=str(invoice.get("supplier") or ""),
            item_ids=tuple(row["id"] for row in invoice_items),
        ))
    if item_exports and any(not item.payment_files for item in item_exports):
        warnings.append(_problem(
            "optional_payment_missing", "部分商品未提供支付凭证；支付凭证为选填，不影响生成。",
        ))
    total = (
        sum((row.total for row in invoice_exports), Decimal("0.00"))
        if invoice_exports and all(row.total is not None for row in invoice_exports)
        else None
    )
    plan = ExportPlan(
        batch_id=str(data.get("id", "")), batch_path=root,
        invoices=tuple(invoice_exports), items=tuple(item_exports),
        sources=tuple(source_map.values()), total=total, warnings=tuple(warnings),
    )
    if errors:
        raise ValidationError(errors, warnings, plan.to_dict())
    return plan


def _verify_outputs(plan, directory):
    expected_names = {row.filename for row in plan.invoices} | {"实物图.docx", "采购统计.xlsx"}
    if {path.name for path in directory.iterdir()} != expected_names:
        raise ValidationError([_problem("output_files", "生成文件数量或名称不一致，未交付本次结果。")])
    sources = {source.id: source for source in plan.sources}
    for invoice in plan.invoices:
        if _sha256(directory / invoice.filename) != sources[invoice.asset_id].sha256:
            raise ValidationError([_problem("output_pdf", "发票复制校验失败，未交付本次结果。")])
    workbook = load_workbook(directory / "采购统计.xlsx", data_only=False)
    try:
        sheet = workbook["采购统计"]
        if workbook.sheetnames != ["采购统计"] or sheet.max_column != 8:
            raise ValueError("工作表名称或列数不正确")
        if [cell.value for cell in sheet[1]] != ExcelGenerator.HEADERS:
            raise ValueError("Excel 表头不正确")
        actual_total = Decimal("0.00")
        for row_index, expected in enumerate(plan.excel_rows, 2):
            actual = [sheet.cell(row_index, col).value for col in range(1, 9)]
            amount = Decimal(str(actual[2]))
            if amount != plan.items[row_index - 2].amount:
                raise ValueError("商品金额不一致")
            if actual[:2] != expected[:2] or actual[3:] != expected[3:]:
                raise ValueError("商品顺序或发票文件引用不一致")
            if any(sheet.cell(row_index, col).data_type == "f" for col in range(1, 9)):
                raise ValueError("商品文本被错误写成公式")
            actual_total += amount
        total_row = len(plan.items) + 3
        if sheet.max_row != total_row or sheet.cell(total_row, 1).value != "总金额":
            raise ValueError("总金额行不正确")
        if Decimal(str(sheet.cell(total_row, 3).value)) != plan.total or actual_total != plan.total:
            raise ValueError("Excel 行合计与发票总金额不一致")
    except (KeyError, ValueError, InvalidOperation, TypeError) as exc:
        raise ValidationError([_problem("output_excel", f"Excel 校验失败：{exc}")]) from exc
    finally:
        workbook.close()
    document = Document(directory / "实物图.docx")
    unique_photos = {path for item in plan.items for path in item.photo_files}
    proof_cards = DocxGenerator._proof_groups(plan.items, explicit_item_numbers=True)
    expected_pictures = len(unique_photos) + sum(
        len(card["proofs"]) + len(card["payments"]) for card in proof_cards
    )
    if len(document.inline_shapes) != expected_pictures:
        raise ValidationError([_problem("output_docx", "Word 图片数量不一致，未交付本次结果。")])


def _commit_directory(staged, target, replace):
    """Swap a completed directory, rolling back if the final rename fails."""
    if target.is_symlink() or (target.exists() and not target.is_dir()):
        raise FileExistsError(f"输出位置已存在其他文件，请选择另一个位置：{target}")
    if target.exists() and not replace:
        raise FileExistsError(f"已存在同金额文件夹，请选择其他位置或明确更新本次归档：{target}")
    backup = None
    if target.exists():
        backup = target.parent / f".agt-previous-{uuid.uuid4().hex}"
        os.replace(target, backup)
    try:
        os.replace(staged, target)
    except BaseException:
        if backup is not None:
            os.replace(backup, target)
        raise
    if backup is not None:
        # The new archive is already complete. A cleanup failure must not report
        # a failed export after successfully replacing the visible destination.
        shutil.rmtree(backup, ignore_errors=True)


def export_plan(
    plan: ExportPlan, output_root: Path, replace=False,
    progress: Callable[[str], None] | None = None,
) -> Path:
    """Generate in a sibling staging directory, validate, then publish the folder."""
    if not isinstance(plan, ExportPlan) or plan.total is None:
        raise ValidationError([_problem("invalid_plan", "请先核对资料并刷新导出预览。")])
    errors = [
        error for source in plan.sources
        for error in _check_source(source, plan.batch_path)
    ]
    if errors:
        raise ValidationError(errors, plan.warnings, plan.to_dict())
    root = Path(output_root).resolve()
    target = root / format(plan.total, ".2f")
    if target.exists() and not replace:
        raise FileExistsError(f"已存在同金额文件夹，请选择其他位置或明确更新本次归档：{target}")
    # Never replace a directory that holds the reviewed source material.
    for source in plan.sources:
        if source.source.is_relative_to(target):
            raise ValidationError([_problem("source_output_overlap", "输出目录包含原始材料，请选择其他输出位置。")])
    root.mkdir(parents=True, exist_ok=True)
    staged = Path(tempfile.mkdtemp(prefix=".agt-export-", dir=root))

    def notify(message):
        if progress:
            progress(message)

    try:
        notify("正在复制原始发票…")
        for invoice in plan.invoices:
            shutil.copy2(invoice.source, staged / invoice.filename)
        notify("正在生成实物图和购买凭证 Word…")
        DocxGenerator().generate(plan.items, staged / "实物图.docx", explicit_item_numbers=True)
        notify("正在生成采购统计 Excel…")
        filename_map = {str(row.source): row.filename for row in plan.invoices}
        ExcelGenerator().generate(
            plan.items, staged / "采购统计.xlsx", plan.total,
            invoice_filename_map=filename_map,
        )
        notify("正在核对发票、Word 与 Excel…")
        _verify_outputs(plan, staged)
        errors = [
            error for source in plan.sources
            for error in _check_source(source, plan.batch_path, inspect_content=False)
        ]
        if errors:
            raise ValidationError(errors, plan.warnings, plan.to_dict())
        _commit_directory(staged, target, replace)
        return target
    finally:
        if staged.exists():
            shutil.rmtree(staged, ignore_errors=True)
