from decimal import Decimal
import hashlib
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import fitz
from docx import Document
from openpyxl import load_workbook
from PIL import Image
import pytest

import core.generators.export_plan as export_module
from core.generators.archive_generator import ArchiveGenerator
from core.generators.docx_generator import DocxGenerator
from core.generators.excel_generator import ExcelGenerator
from core.generators.export_plan import ValidationError, build_export_plan, export_plan
from core.graph.purchase_graph import InvoiceDocument, PurchaseGraph, PurchaseItem


def _asset(batch, data, asset_id, kind, *, item_ids=(), confirmed=True):
    suffix = ".pdf" if kind == "invoice" else ".png"
    path = batch / "assets" / f"{asset_id}{suffix}"
    path.parent.mkdir(parents=True, exist_ok=True)
    if kind == "invoice":
        with fitz.open() as pdf:
            page = pdf.new_page()
            page.insert_text((50, 50), f"Original invoice {asset_id}")
            pdf.save(path)
    else:
        color = hashlib.sha256(asset_id.encode()).digest()[:3]
        Image.new("RGB", (240, 400), tuple(color)).save(path)
    row = {
        "id": asset_id, "path": path.relative_to(batch).as_posix(),
        "original_name": f"乱码-未命名-{asset_id}{suffix}",
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "kind": kind, "item_ids": list(item_ids),
        "confirmed": confirmed, "ignored": False, "error": None,
    }
    data["assets"].append(row)
    return row


def _batch(tmp_path):
    batch = tmp_path / "batch"
    batch.mkdir()
    data = {"id": "batch-1", "name": "采购测试", "assets": [], "invoices": [], "items": [], "exports": []}
    _asset(batch, data, "pdf_a", "invoice")
    _asset(batch, data, "pdf_b", "invoice")
    # The authoritative invoice order is deliberately different from names,
    # the import sequence, and the flat item list.
    data["invoices"] = [
        {"id": "inv_b", "asset_id": "pdf_b", "total": "30.00", "invoice_no": "002", "confirmed": True},
        {"id": "inv_a", "asset_id": "pdf_a", "total": "30.00", "invoice_no": "001", "confirmed": True},
    ]
    data["items"] = [
        {"id": "a1", "invoice_id": "inv_a", "name": "同名商品", "raw_name": "原名-A1", "amount": "10.00"},
        {"id": "a2", "invoice_id": "inv_a", "name": "同名商品", "raw_name": "原名-A2", "amount": "20.00"},
        {"id": "b1", "invoice_id": "inv_b", "name": "商品B", "raw_name": "原名-B", "amount": "30.00"},
    ]
    return data, batch


def _codes(error):
    return {row["code"] for row in error.value.errors}


def test_authoritative_invoice_order_shared_photo_and_eight_column_export(tmp_path):
    data, batch = _batch(tmp_path)
    _asset(batch, data, "shared_photo", "photo", item_ids=("a1", "b1"))
    _asset(batch, data, "second_photo", "photo", item_ids=("a2",))
    _asset(batch, data, "order_a", "purchase_proof", item_ids=("a1", "a2"))
    _asset(batch, data, "order_b", "purchase_proof", item_ids=("b1",))
    _asset(batch, data, "payment_b", "payment_proof", item_ids=("b1",))
    plan = build_export_plan(data, batch)
    assert [item.id for item in plan.items] == ["b1", "a1", "a2"]
    assert plan.total == Decimal("60.00")
    assert plan.items[1].photo_ids == ("shared_photo",)
    assert plan.items[2].photo_ids == ("second_photo",)  # Same names cannot merge IDs.
    assert plan.items[1].invoice_filename == plan.items[2].invoice_filename
    assert plan.to_dict()["counts"] == {
        "invoices": 2, "items": 3, "photos": 2, "purchase_proofs": 2, "payment_proofs": 1,
    }
    output = export_plan(plan, tmp_path / "output")
    assert output.name == "60.00"
    expected_files = {row.filename for row in plan.invoices} | {"实物图.docx", "采购统计.xlsx"}
    assert {path.name for path in output.iterdir()} == expected_files
    for invoice in plan.invoices:
        assert (output / invoice.filename).read_bytes() == invoice.source.read_bytes()
    wb = load_workbook(output / "采购统计.xlsx")
    ws = wb["采购统计"]
    assert [cell.value for cell in ws[1]] == ExcelGenerator.HEADERS
    assert ws.max_column == 8
    assert [ws.cell(row, 1).value for row in range(2, 5)] == [1, 2, 3]
    assert [ws.cell(row, 3).value for row in range(2, 5)] == [30, 10, 20]
    assert ws.cell(6, 1).value == "总金额"
    assert ws.cell(6, 3).value == 60
    for row, item in enumerate(plan.items, 2):
        assert ws.cell(row, 8).value == item.invoice_filename
        assert (output / ws.cell(row, 8).value).is_file()
    wb.close()
    doc = Document(output / "实物图.docx")
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "对应商品：01 商品B + 02 同名商品" in text
    assert "对应商品：03 同名商品" in text
    assert len(doc.inline_shapes) == 5  # Shared photo once, another photo, two orders, payment.
    assert len(doc.tables) == 1 and len(doc.tables[0].rows[0].cells) == 3


def test_missing_optional_materials_still_export_and_does_not_count_boxes(tmp_path):
    data, batch = _batch(tmp_path)
    _asset(batch, data, "box", "package", item_ids=("a1",))
    _asset(batch, data, "uncertain", "unknown", item_ids=("a1",))
    _asset(batch, data, "suggestion", "photo", item_ids=("a1",), confirmed=False)
    plan = build_export_plan(data, batch)
    assert all(not item.photo_files for item in plan.items)
    assert "optional_payment_missing" in {warning["code"] for warning in plan.warnings}
    assert export_plan(plan, tmp_path / "output").is_dir()


def test_partial_draft_errors_include_preview_and_zero_is_a_valid_amount(tmp_path):
    data, batch = _batch(tmp_path)
    data["invoices"][0]["confirmed"] = False
    data["items"][2]["amount"] = None
    with pytest.raises(ValidationError) as error:
        build_export_plan(data, batch)
    assert {"unconfirmed_invoice", "missing_amount"} <= _codes(error)
    assert error.value.preview["counts"]["items"] == 3
    assert error.value.preview["items"][0]["amount"] is None
    data["invoices"][0].update(confirmed=True, total="0.00")
    data["items"][2]["amount"] = "0.00"
    assert build_export_plan(data, batch).total == Decimal("30.00")


@pytest.mark.parametrize("amount", [None, "", "NaN", "Infinity", "-Infinity", "1.001"])
def test_invalid_money_is_not_coerced_to_zero(tmp_path, amount):
    data, batch = _batch(tmp_path)
    data["items"][0]["amount"] = amount
    with pytest.raises(ValidationError):
        build_export_plan(data, batch)


def test_equal_batch_total_cannot_hide_opposite_per_invoice_differences(tmp_path):
    data, batch = _batch(tmp_path)
    data["items"][0]["amount"] = "11.00"
    data["items"][2]["amount"] = "29.00"
    assert sum(Decimal(item["amount"]) for item in data["items"]) == Decimal("60.00")
    with pytest.raises(ValidationError) as error:
        build_export_plan(data, batch)
    mismatches = [row for row in error.value.errors if row["code"] == "invoice_item_amount_mismatch"]
    assert {row["invoice_id"] for row in mismatches} == {"inv_a", "inv_b"}


def test_duplicate_invoice_numbers_are_blocked_and_ignored_invoice_can_be_restored(tmp_path):
    data, batch = _batch(tmp_path)
    data["invoices"][0]["invoice_no"] = "001"
    with pytest.raises(ValidationError) as error:
        build_export_plan(data, batch)
    assert "duplicate_invoice" in _codes(error)
    data["assets"][0]["ignored"] = True
    plan = build_export_plan(data, batch)
    assert len(plan.invoices) == 1 and [item.id for item in plan.items] == ["b1"]
    assert len(data["items"]) == 3
    data["assets"][0]["ignored"] = False
    with pytest.raises(ValidationError):
        build_export_plan(data, batch)


def test_duplicate_pdf_bytes_are_blocked_without_invoice_numbers(tmp_path):
    data, batch = _batch(tmp_path)
    first, second = data["assets"]
    (batch / second["path"]).write_bytes((batch / first["path"]).read_bytes())
    second["sha256"] = first["sha256"]
    for invoice in data["invoices"]:
        invoice["invoice_no"] = ""
    with pytest.raises(ValidationError) as error:
        build_export_plan(data, batch)
    assert "duplicate_invoice_file" in _codes(error)


@pytest.mark.parametrize("change", ["missing", "changed"])
def test_sources_are_rechecked_between_preview_and_export(tmp_path, change):
    data, batch = _batch(tmp_path)
    plan = build_export_plan(data, batch)
    source = plan.invoices[0].source
    if change == "missing":
        source.unlink()
    else:
        source.write_bytes(source.read_bytes() + b"\nmodified")
    output = tmp_path / "output"
    with pytest.raises(ValidationError):
        export_plan(plan, output)
    assert not output.exists()


@pytest.mark.parametrize("relative_path", ["../invoice.pdf", "/tmp/invoice.pdf", r"C:\outside\invoice.pdf"])
def test_batch_paths_cannot_escape_the_batch(tmp_path, relative_path):
    data, batch = _batch(tmp_path)
    data["assets"][0]["path"] = relative_path
    with pytest.raises(ValidationError) as error:
        build_export_plan(data, batch)
    assert "invalid_source" in _codes(error)


def test_symlink_cannot_escape_batch(tmp_path):
    data, batch = _batch(tmp_path)
    source = batch / data["assets"][0]["path"]
    outside = tmp_path / "outside.pdf"
    source.replace(outside)
    try:
        source.symlink_to(outside)
    except OSError:
        pytest.skip("Symlink creation is not permitted on this host")
    with pytest.raises(ValidationError):
        build_export_plan(data, batch)


def test_output_collision_and_explicit_update_leave_no_stale_pdfs(tmp_path):
    data, batch = _batch(tmp_path)
    first = build_export_plan(data, batch)
    output = export_plan(first, tmp_path / "output")
    previous_name = first.invoices[0].filename
    data["items"][2]["name"] = "更新后的商品名"
    second = build_export_plan(data, batch)
    with pytest.raises(FileExistsError):
        export_plan(second, output.parent)
    assert (output / previous_name).is_file()
    (output / "旧文件.pdf").write_bytes(b"stale")
    updated = export_plan(second, output.parent, replace=True)
    assert updated == output
    assert not (updated / previous_name).exists()
    assert not (updated / "旧文件.pdf").exists()
    assert {p.name for p in updated.iterdir()} == (
        {row.filename for row in second.invoices} | {"实物图.docx", "采购统计.xlsx"}
    )


def test_failed_update_keeps_previous_folder_and_removes_temporary_output(tmp_path, monkeypatch):
    data, batch = _batch(tmp_path)
    plan = build_export_plan(data, batch)
    output = export_plan(plan, tmp_path / "output")
    before = {p.name: p.read_bytes() for p in output.iterdir()}

    def fail(*args, **kwargs):
        raise RuntimeError("simulated generation error")

    monkeypatch.setattr(DocxGenerator, "generate", fail)
    with pytest.raises(RuntimeError):
        export_plan(plan, output.parent, replace=True)
    assert {p.name: p.read_bytes() for p in output.iterdir()} == before
    assert list(output.parent.iterdir()) == [output]


def test_final_swap_failure_rolls_back_previous_folder(tmp_path, monkeypatch):
    data, batch = _batch(tmp_path)
    plan = build_export_plan(data, batch)
    output = export_plan(plan, tmp_path / "output")
    before = {p.name: p.read_bytes() for p in output.iterdir()}
    real_replace = export_module.os.replace
    failed = False

    def fail_once(source, target):
        nonlocal failed
        if Path(source).name.startswith(".agt-export-") and not failed:
            failed = True
            raise OSError("simulated final rename failure")
        return real_replace(source, target)

    monkeypatch.setattr(export_module.os, "replace", fail_once)
    with pytest.raises(OSError):
        export_plan(plan, output.parent, replace=True)
    assert {p.name: p.read_bytes() for p in output.iterdir()} == before
    assert list(output.parent.iterdir()) == [output]


def test_text_starting_with_equals_remains_literal_in_unchanged_excel_format(tmp_path):
    data, batch = _batch(tmp_path)
    data["items"][0]["name"] = '=HYPERLINK("https://example.invalid","商品")'
    plan = build_export_plan(data, batch)
    output = export_plan(plan, tmp_path / "output")
    wb = load_workbook(output / "采购统计.xlsx", data_only=False)
    cell = wb["采购统计"].cell(3, 2)
    assert cell.value == data["items"][0]["name"]
    assert cell.data_type == "s"
    wb.close()


def test_invalid_generated_workbook_is_not_published(tmp_path, monkeypatch):
    data, batch = _batch(tmp_path)
    plan = build_export_plan(data, batch)
    original = ExcelGenerator.generate

    def incorrect(self, items, output, total, **kwargs):
        original(self, items, output, total, **kwargs)
        wb = load_workbook(output)
        wb["采购统计"].cell(2, 8).value = "原乱码文件.pdf"
        wb.save(output)
        wb.close()

    monkeypatch.setattr(ExcelGenerator, "generate", incorrect)
    with pytest.raises(ValidationError) as error:
        export_plan(plan, tmp_path / "output")
    assert "output_excel" in _codes(error)
    assert not (tmp_path / "output" / "60.00").exists()


def test_parser_error_can_be_repaired_manually_without_rerunning_parser(tmp_path):
    data, batch = _batch(tmp_path)
    data["assets"][0]["error"] = "OCR not available, table could not be recognized"
    assert build_export_plan(data, batch).total == Decimal("60.00")


def test_an_invoice_cannot_reuse_a_verified_photo_asset_as_a_pdf(tmp_path):
    data, batch = _batch(tmp_path)
    photo = _asset(batch, data, "photo", "photo", item_ids=("a1",))
    data["invoices"][0]["asset_id"] = photo["id"]
    with pytest.raises(ValidationError) as error:
        build_export_plan(data, batch)
    assert "wrong_material_type" in _codes(error)


def test_phone_photo_orientation_and_webp_are_embedded_as_readable_pictures(tmp_path):
    data, batch = _batch(tmp_path)
    phone = _asset(batch, data, "phone", "photo", item_ids=("a1",))
    phone_path = batch / "assets" / "phone.jpg"
    exif = Image.Exif()
    exif[274] = 6
    Image.new("RGB", (100, 200), "green").save(phone_path, exif=exif)
    phone["path"] = phone_path.relative_to(batch).as_posix()
    phone["sha256"] = hashlib.sha256(phone_path.read_bytes()).hexdigest()
    webp = _asset(batch, data, "webp", "photo", item_ids=("b1",))
    webp_path = batch / "assets" / "photo.webp"
    Image.new("RGB", (300, 400), "blue").save(webp_path)
    webp["path"] = webp_path.relative_to(batch).as_posix()
    webp["sha256"] = hashlib.sha256(webp_path.read_bytes()).hexdigest()
    output = export_plan(build_export_plan(data, batch), tmp_path / "output")
    with ZipFile(output / "实物图.docx") as archive:
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
        sizes = []
        for name in media:
            with Image.open(BytesIO(archive.read(name))) as image:
                sizes.append(image.size)
        assert sorted(sizes) == [(200, 100), (300, 400)]


def test_legacy_generator_does_not_silently_skip_a_missing_pdf(tmp_path):
    source = tmp_path / "missing.pdf"
    graph = PurchaseGraph(
        items=[PurchaseItem(name="商品", amount=10, invoice_files=[str(source)])],
        invoices=[InvoiceDocument(path=str(source), total=10, item_names=["商品"])],
        total_amount=10,
    )
    with pytest.raises(ValidationError):
        ArchiveGenerator().generate(graph, tmp_path / "output")
    assert not (tmp_path / "output").exists()


def test_legacy_generator_respects_failed_consistency_check(tmp_path):
    source = tmp_path / "source.pdf"
    source.write_bytes(b"%PDF-original")
    graph = PurchaseGraph(
        items=[PurchaseItem(name="商品", amount=9, invoice_files=[str(source)])],
        invoices=[InvoiceDocument(path=str(source), total=10, item_names=["商品"])],
        total_amount=10,
    )
    with pytest.raises(ValidationError):
        ArchiveGenerator().generate(graph, tmp_path / "output")
    assert not (tmp_path / "output").exists()
